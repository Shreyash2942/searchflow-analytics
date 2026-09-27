"""Build the APA 7 student-paper DOCX from apa_project_report.md."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(__file__).with_name("apa_project_report.md")
OUTPUT = Path(__file__).with_name("APA_SearchFlow_Analytics_Report.docx")


def _clean_inline(value: str) -> str:
    """Remove Markdown-only syntax while keeping readable report text."""
    value = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", value)
    value = value.replace("`", "").replace("**", "").replace("*", "")
    return value


def _add_page_number(section) -> None:
    """Place an APA-style page number in the top-right header."""
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.append(field)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def _set_cell_border(cell, **kwargs) -> None:
    """Apply table-border attributes to one cell."""
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            tag = "w:" + edge
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            for key in ["val", "sz", "space", "color"]:
                if key in kwargs[edge]:
                    element.set(qn("w:" + key), str(kwargs[edge][key]))


def _add_table(document: Document, rows: list[list[str]]) -> None:
    """Add a simple APA-style table with no vertical borders."""
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = True
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = _clean_inline(value.strip())
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)
                    if row_index == 0:
                        run.bold = True
            _set_cell_border(
                cell,
                top={"val": "single", "sz": 6, "color": "000000"}
                if row_index == 0 else {"val": "nil"},
                bottom={"val": "single", "sz": 6, "color": "000000"}
                if row_index in (0, len(rows) - 1) else {"val": "nil"},
                left={"val": "nil"},
                right={"val": "nil"},
                insideV={"val": "nil"},
            )
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _add_body_paragraph(document: Document, text: str, *, indent: bool = True) -> None:
    """Add an APA double-spaced body paragraph."""
    paragraph = document.add_paragraph(_clean_inline(text))
    paragraph.paragraph_format.line_spacing = 2
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.first_line_indent = Inches(0.5) if indent else Inches(0)


def _add_heading(document: Document, text: str, level: int) -> None:
    """Add a centered APA Level 1 or left-aligned Level 2 heading."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 2
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(_clean_inline(text))
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def _title_page(document: Document, lines: list[str]) -> None:
    """Create the APA student-paper title page."""
    for _ in range(5):
        document.add_paragraph()
    for index, line in enumerate(lines):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing = 2
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(_clean_inline(line).strip())
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        if index == 0:
            run.bold = True


def build() -> Path:
    """Parse the Markdown source and write a formatted APA 7 DOCX."""
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    _add_page_number(section)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)

    title_end = lines.index("## Abstract")
    title_lines = [_clean_inline(lines[0].lstrip("#")).strip()]
    title_lines.extend(_clean_inline(line).strip() for line in lines[1:title_end] if line.strip())
    _title_page(document, title_lines)
    document.add_page_break()

    index = title_end
    paragraph_buffer: list[str] = []
    in_code = False
    in_table = False
    table_rows: list[list[str]] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            text = " ".join(part.strip() for part in paragraph_buffer)
            _add_body_paragraph(document, text)
            paragraph_buffer = []

    def flush_table() -> None:
        nonlocal table_rows
        if table_rows:
            _add_table(document, table_rows)
            table_rows = []

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if stripped.startswith("```"):
            flush_paragraph()
            flush_table()
            in_code = not in_code
            index += 1
            continue
        if in_code:
            paragraph = document.add_paragraph(_clean_inline(raw))
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.name = "Courier New"
                run.font.size = Pt(10)
            index += 1
            continue
        if stripped.startswith("## "):
            flush_paragraph()
            flush_table()
            heading = stripped[3:]
            if heading == "Introduction":
                document.add_page_break()
                title_paragraph = document.add_paragraph()
                title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                title_paragraph.paragraph_format.line_spacing = 2
                title_paragraph.paragraph_format.space_after = Pt(0)
                title_run = title_paragraph.add_run(title_lines[0])
                title_run.bold = True
                title_run.font.name = "Times New Roman"
                title_run.font.size = Pt(12)
            if heading.startswith("Appendix"):
                document.add_page_break()
            _add_heading(document, heading, 1)
            index += 1
            continue
        if stripped.startswith("### "):
            flush_paragraph()
            flush_table()
            _add_heading(document, stripped[4:], 2)
            index += 1
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            flush_paragraph()
            values = [part.strip() for part in stripped.strip("|").split("|")]
            if not all(set(value) <= set("-:") for value in values):
                table_rows.append(values)
            in_table = True
            index += 1
            continue
        if in_table:
            flush_table()
            in_table = False
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        if stripped.startswith("- "):
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            paragraph.paragraph_format.line_spacing = 2
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(_clean_inline(stripped[2:]))
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            index += 1
            continue
        if stripped.startswith("*Keywords:*"):
            flush_paragraph()
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.line_spacing = 2
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run("Keywords:")
            run.italic = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run = paragraph.add_run(_clean_inline(stripped[len("*Keywords:*"):]))
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            index += 1
            continue
        if stripped.startswith("**Table "):
            flush_paragraph()
            paragraph = document.add_paragraph(_clean_inline(stripped))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.line_spacing = 2
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
            index += 1
            continue
        if heading := re.match(r"\*([^*]+)\*", stripped):
            flush_paragraph()
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.line_spacing = 2
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(heading.group(1))
            run.italic = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            index += 1
            continue
        paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph()
    flush_table()

    # References use hanging indents in APA style.
    reference_index = None
    for candidate_index, paragraph in enumerate(document.paragraphs):
        if paragraph.text.strip() == "References":
            references_heading = paragraph
            reference_index = candidate_index
            break
    else:
        references_heading = None
    if references_heading is not None:
        start = reference_index + 1
        for paragraph in document.paragraphs[start:]:
            if paragraph.text.strip().startswith("Appendix"):
                break
            if paragraph.text.strip():
                paragraph.paragraph_format.left_indent = Inches(0.5)
                paragraph.paragraph_format.first_line_indent = Inches(-0.5)

    document.core_properties.title = "SearchFlow Analytics: A Measured Comparison of Linear and Binary Search"
    document.core_properties.author = "Shreyash2942"
    document.core_properties.subject = "APA 7 portfolio report"
    document.core_properties.keywords = "linear search, binary search, algorithms, benchmarking"
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
